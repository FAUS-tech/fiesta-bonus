"""Backtest the boss's proposed comp plan against the current plan
under multiple behavior-change scenarios."""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from build_bonus_workbook import (
    AGENTS, AGENT_NAMES, MONTHS, current_bonus,
    TITLE_FONT, SECTION_FONT, SECTION_FILL, SUB_FILL, CURRENT_FILL, PROP_A_FILL,
    style_header, style_data, style_dollar, style_int, set_col_widths,
    swap_rwr_to_ren, full_flip,
)

# ============================================================================
# BOSS'S PROPOSED PLAN
# ============================================================================
PROPOSAL_TIERS = {
    'Trainee':    {'base': 34000, 'nb_pct': 0.05, 'rn_pct': 0.000, 'rwr_pct': 0.00, 'min_nb_count': 0},
    'T1 Builder': {'base': 36000, 'nb_pct': 0.08, 'rn_pct': 0.015, 'rwr_pct': 0.00, 'min_nb_count': 25},
    'T2 Senior':  {'base': 40000, 'nb_pct': 0.12, 'rn_pct': 0.025, 'rwr_pct': 0.02, 'min_nb_count': 50},
    'T3 Elite':   {'base': 45000, 'nb_pct': 0.16, 'rn_pct': 0.035, 'rwr_pct': 0.02, 'min_nb_count': 80},
    'Elite Top':  {'base': 50000, 'nb_pct': 0.20, 'rn_pct': 0.050, 'rwr_pct': 0.02, 'min_nb_count': 110},
}
NB_GATE = 45000
CARRIER_NB_COMM = 0.10
CARRIER_REN_COMM = 0.08
CARRIER_RWR_COMM = 0.08


def collected_kicker_per_policy(coll_pct):
    if coll_pct >= 1.00: return 8
    if coll_pct >= 0.25: return 5
    if coll_pct >= 0.20: return 2
    if coll_pct >= 0.15: return 1
    return 0


def assign_tier(avg_nb):
    if avg_nb >= 110: return 'Elite Top'
    if avg_nb >=  80: return 'T3 Elite'
    if avg_nb >=  50: return 'T2 Senior'
    if avg_nb >=  25: return 'T1 Builder'
    return 'Trainee'


def grow_nb(d, factor):
    """Multiply NB count, premium, collected by 'factor' (1.25 = +25% growth)."""
    nb_c, nb_p, nb_col = d['NB']
    return {
        'NB':  (int(round(nb_c * factor)), nb_p * factor, nb_col * factor),
        'RWR': d['RWR'],
        'REN': d['REN'],
    }


def chain(*fns):
    """Compose scenario transforms left-to-right."""
    def inner(d):
        for f in fns:
            d = f(d)
        return d
    return inner


# ============================================================================
# SCENARIOS
# ============================================================================
# Each scenario is a callable that takes a raw agent-month dict and returns
# a transformed one. The transforms compound.
SCENARIOS = [
    ('REAL (today)',                   'No behavior change. Actual Jan-Apr 2026 data.',
        lambda d: d),
    ('SWAP 50% (half RWR -> REN)',     'Agents retain instead of rewriting half the book.',
        lambda d: swap_rwr_to_ren(d, 0.5)),
    ('FLIP 100% (all RWR -> REN)',     'Best-case retention behavior shift.',
        full_flip),
    ('GROWTH +25% NB + FLIP',          'Agents write 25% more NB and stop rewriting entirely.',
        chain(full_flip, lambda d: grow_nb(d, 1.25))),
    ('GROWTH +50% NB + FLIP',          'Agents write 50% more NB and stop rewriting entirely.',
        chain(full_flip, lambda d: grow_nb(d, 1.50))),
]


def compute_proposal_bonus(agent_data, tier):
    """Compute the boss-proposed monthly bonus for one agent-month."""
    t = PROPOSAL_TIERS[tier]
    nb_c, nb_p, nb_col = agent_data['NB']
    rwr_c, rwr_p, _ = agent_data['RWR']
    ren_c, ren_p, _ = agent_data['REN']
    gate_met = nb_p >= NB_GATE
    if not gate_met:
        return {'gate_met': False, 'nb_var': 0, 'rn_var': 0, 'rwr_var': 0,
                'kicker': 0, 'total': 0}
    nb_var = nb_p * CARRIER_NB_COMM * t['nb_pct']
    rn_var = ren_p * CARRIER_REN_COMM * t['rn_pct']
    rwr_var = rwr_p * CARRIER_RWR_COMM * t['rwr_pct']
    nb_coll_pct = nb_col / nb_p if nb_p else 0
    kicker = nb_c * collected_kicker_per_policy(nb_coll_pct)
    return {'gate_met': True, 'nb_var': nb_var, 'rn_var': rn_var,
            'rwr_var': rwr_var, 'kicker': kicker,
            'total': nb_var + rn_var + rwr_var + kicker}


def compute_scenario(scenario_fn):
    """Run one scenario across all 6 agents x 4 months. Returns:
       - per_agent: {agent: {'cur', 'prop', 'tier', 'avg_nb', 'gate_pass_months'}}
       - tot_cur, tot_prop, gate_pass_count_total
    """
    # First pass: compute each agent's avg NB in the scenario to pick tier
    per_agent = {}
    for agent in AGENT_NAMES:
        nb_sum = 0
        for month in MONTHS:
            d = scenario_fn(AGENTS[agent][month])
            nb_sum += d['NB'][0]
        avg_nb = nb_sum / 4
        per_agent[agent] = {
            'avg_nb': avg_nb,
            'tier': assign_tier(avg_nb),
            'cur': 0,
            'prop': 0,
            'gate_pass_months': 0,
            'monthly': [],   # list of dicts per month for detail rendering
        }

    # Second pass: monthly numbers
    tot_cur = tot_prop = gate_pass_total = 0
    for agent in AGENT_NAMES:
        tier = per_agent[agent]['tier']
        for month in MONTHS:
            d = scenario_fn(AGENTS[agent][month])
            res = compute_proposal_bonus(d, tier)
            cur = current_bonus(d['NB'][0], d['RWR'][0])
            per_agent[agent]['cur'] += cur
            per_agent[agent]['prop'] += res['total']
            if res['gate_met']:
                per_agent[agent]['gate_pass_months'] += 1
                gate_pass_total += 1
            per_agent[agent]['monthly'].append({
                'month': month, 'd': d, 'res': res, 'cur': cur,
            })
            tot_cur += cur
            tot_prop += res['total']

    return per_agent, tot_cur, tot_prop, gate_pass_total


# ============================================================================
# RENDERING HELPERS
# ============================================================================
GREEN_FILL = PatternFill('solid', fgColor='00B050')
RED_FILL = PatternFill('solid', fgColor='C00000')
GREY_FILL = PatternFill('solid', fgColor='D9D9D9')
GATE_PASS_FILL = PatternFill('solid', fgColor='C6EFCE')
GATE_MISS_FILL = PatternFill('solid', fgColor='FFC7CE')
NAVY_FILL = PatternFill('solid', fgColor='1F4E78')


def build_scenarios_summary(wb, scenario_results):
    ws = wb.create_sheet('Scenarios Summary')
    ws['A1'] = "Boss Proposal vs Current - 5 Behavior Scenarios"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:G1')
    ws['A2'] = ("Same 6 agents x 4 months (Jan-Apr 2026). Each scenario applies a behavior transform "
                "(NB growth, RWR -> REN flip) and reruns the boss-proposed plan against today's "
                "count-tier plan. Tiers re-assigned per scenario based on the scenario's avg NB.")
    ws['A2'].font = Font(italic=True, size=10, color='666666')
    ws.merge_cells('A2:G2')

    r = 4
    ws.cell(row=r, column=1, value='HEADLINE - TOTAL BONUS PAID (4-month, 6 agents)').font = SECTION_FONT
    ws.cell(row=r, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    r += 1

    headers = ['Scenario', 'Description', 'Current Plan', 'Boss Proposal',
               'Delta (proposal - current)', 'Gate pass (of 24)', 'Annualized proposal']
    for i, h in enumerate(headers, 1):
        style_header(ws.cell(row=r, column=i, value=h))
    ws.row_dimensions[r].height = 28
    r += 1

    for (label, desc, _), result in zip(SCENARIOS, scenario_results):
        _, tot_cur, tot_prop, gate_pass = result
        delta = tot_prop - tot_cur
        vals = [label, desc, tot_cur, tot_prop, delta,
                f"{gate_pass}/24 ({gate_pass/24*100:.0f}%)", tot_prop * 3]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            if i in (1,): style_data(c); c.font = Font(bold=True)
            elif i in (2, 6): style_data(c)
            else: style_dollar(c)
        # Color the delta cell
        delta_cell = ws.cell(row=r, column=5)
        delta_cell.font = Font(bold=True, color='FFFFFF')
        delta_cell.fill = GREEN_FILL if delta > 0 else (RED_FILL if delta < 0 else GREY_FILL)
        delta_cell.alignment = Alignment(horizontal='center')
        ws.cell(row=r, column=3).fill = CURRENT_FILL
        ws.cell(row=r, column=4).fill = PROP_A_FILL
        ws.row_dimensions[r].height = 28
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='READ').font = SECTION_FONT
    ws.cell(row=r, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    r += 1
    real = scenario_results[0]
    flip = scenario_results[2]
    growth50 = scenario_results[4]
    reads = [
        f"REAL behavior: boss proposal pays ${real[2]:,.0f} vs today's ${real[1]:,.0f} - a {(real[2]-real[1])/real[1]*100:.0f}% drop. The $45K NB gate is the bottleneck (only {real[3]} of 24 months pass).",
        f"FLIP scenario (no rewrites, all retained as renewals): proposal pays ${flip[2]:,.0f}. Almost unchanged from REAL because the gate is on NB premium - flipping RWR to REN doesn't unlock the gate.",
        f"GROWTH +50% NB scenario: proposal pays ${growth50[2]:,.0f} ({growth50[2]/real[1]*100:.0f}% of today's pay). NOW the gate opens for more agents - {growth50[3]} of 24 months pass.",
        f"Bottom line: the boss proposal only rewards bigger NB books. Agent behavior shifts toward retention alone don't move the needle - they have to grow NB to break the gate.",
        f"Compare: our designed plan rewards retention through the REN tier ladder directly, so SWAP/FLIP scenarios pay more even without NB growth.",
    ]
    for txt in reads:
        c = ws.cell(row=r, column=1, value=txt)
        c.font = Font(size=11)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 44
        r += 1

    set_col_widths(ws, [28, 42, 14, 14, 17, 16, 16])


def build_per_agent_grid(wb, scenario_results):
    """One row per agent. Columns = scenarios. Cell = boss proposal $ for that
    agent in that scenario. Adds current-plan column at far right."""
    ws = wb.create_sheet('Per-Agent Grid')
    ws['A1'] = "Per-Agent View - Boss Proposal Across Scenarios"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')
    ws['A2'] = "Each cell = the boss-proposed 4-month bonus for that agent under that scenario. Last column = today's actual plan for the same agent."
    ws['A2'].font = Font(italic=True, size=10, color='666666')
    ws.merge_cells('A2:H2')

    r = 4
    headers = ['Agent'] + [s[0].split(' (')[0].split('  ')[0][:20] for s in SCENARIOS] + ['Current (today)']
    for i, h in enumerate(headers, 1):
        style_header(ws.cell(row=r, column=i, value=h))
    ws.row_dimensions[r].height = 32
    r += 1

    for agent in AGENT_NAMES:
        ws.cell(row=r, column=1, value=agent).font = Font(bold=True)
        style_data(ws.cell(row=r, column=1))
        col = 2
        for result in scenario_results:
            per_agent, _, _, _ = result
            c = ws.cell(row=r, column=col, value=per_agent[agent]['prop'])
            style_dollar(c); c.fill = PROP_A_FILL
            col += 1
        c = ws.cell(row=r, column=col, value=scenario_results[0][0][agent]['cur'])
        style_dollar(c); c.fill = CURRENT_FILL; c.font = Font(bold=True)
        r += 1

    # Totals row
    ws.cell(row=r, column=1, value='TOTAL').font = Font(bold=True)
    col = 2
    for result in scenario_results:
        _, _, tot_prop, _ = result
        c = ws.cell(row=r, column=col, value=tot_prop)
        style_dollar(c); c.font = Font(bold=True, color='FFFFFF'); c.fill = NAVY_FILL
        c.alignment = Alignment(horizontal='center')
        col += 1
    c = ws.cell(row=r, column=col, value=scenario_results[0][1])
    style_dollar(c); c.font = Font(bold=True, color='FFFFFF'); c.fill = NAVY_FILL
    c.alignment = Alignment(horizontal='center')
    for k in range(1, col + 1):
        if not ws.cell(row=r, column=k).fill or ws.cell(row=r, column=k).fill.fgColor.rgb == '00000000':
            ws.cell(row=r, column=k).fill = SUB_FILL
    ws.row_dimensions[r].height = 26

    set_col_widths(ws, [24, 16, 16, 16, 16, 16, 16, 16])


def build_tier_assignment_grid(wb, scenario_results):
    """Shows how the auto-assigned tier changes by scenario. Useful for
    visualizing how growth pushes agents up the tier ladder."""
    ws = wb.create_sheet('Tier Assignments by Scenario')
    ws['A1'] = "Tier Auto-Assignment Across Scenarios"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:G1')
    ws['A2'] = "Tier is assigned by the scenario's 4-month avg NB count. Growth scenarios push more agents into T1/T2."
    ws['A2'].font = Font(italic=True, size=10, color='666666')
    ws.merge_cells('A2:G2')

    r = 4
    headers = ['Agent'] + [s[0].split(' (')[0][:20] for s in SCENARIOS]
    for i, h in enumerate(headers, 1):
        style_header(ws.cell(row=r, column=i, value=h))
    ws.row_dimensions[r].height = 32
    r += 1

    tier_colors = {
        'Trainee':    PatternFill('solid', fgColor='F4B084'),
        'T1 Builder': PatternFill('solid', fgColor='FFD966'),
        'T2 Senior':  PatternFill('solid', fgColor='92D050'),
        'T3 Elite':   PatternFill('solid', fgColor='00B050'),
        'Elite Top':  PatternFill('solid', fgColor='1F4E78'),
    }

    for agent in AGENT_NAMES:
        ws.cell(row=r, column=1, value=agent).font = Font(bold=True)
        style_data(ws.cell(row=r, column=1))
        col = 2
        for result in scenario_results:
            per_agent, _, _, _ = result
            tier = per_agent[agent]['tier']
            avg_nb = per_agent[agent]['avg_nb']
            text_color = 'FFFFFF' if tier == 'Elite Top' else '000000'
            c = ws.cell(row=r, column=col, value=f"{tier}\n({avg_nb:.0f} NB/mo)")
            c.fill = tier_colors.get(tier, PatternFill())
            c.font = Font(bold=True, size=10, color=text_color)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            col += 1
        ws.row_dimensions[r].height = 36
        r += 1

    set_col_widths(ws, [24, 18, 18, 18, 18, 18, 18])


def build_month_detail(wb, scenario_results, scenario_idx, sheet_name):
    """Stack the month-level detail for one scenario."""
    ws = wb.create_sheet(sheet_name)
    label, desc, _ = SCENARIOS[scenario_idx]
    ws['A1'] = f"{label} - Month by Month Detail"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:N1')
    ws['A2'] = desc
    ws['A2'].font = Font(italic=True, size=10, color='666666')
    ws.merge_cells('A2:N2')

    per_agent, tot_cur, tot_prop, gate_pass = scenario_results[scenario_idx]

    r = 4
    headers = ['Agent', 'Tier', 'Month',
               'NB #', 'NB $', 'NB Coll %',
               'RWR #', 'REN $',
               '$45K Gate', 'NB var', 'Kicker',
               'BOSS Proposal', 'CURRENT', 'Delta']
    for i, h in enumerate(headers, 1):
        style_header(ws.cell(row=r, column=i, value=h))
    ws.row_dimensions[r].height = 26
    r += 1

    for agent in AGENT_NAMES:
        info = per_agent[agent]
        tier = info['tier']
        for m_info in info['monthly']:
            d = m_info['d']
            res = m_info['res']
            cur = m_info['cur']
            nb_c, nb_p, nb_col = d['NB']
            nb_coll_pct = nb_col / nb_p if nb_p else 0
            rwr_c = d['RWR'][0]
            ren_p = d['REN'][1]
            delta = res['total'] - cur

            vals = [agent, tier, m_info['month'], nb_c, nb_p, nb_coll_pct,
                    rwr_c, ren_p,
                    'PASS' if res['gate_met'] else 'MISS',
                    res['nb_var'] + res['rn_var'] + res['rwr_var'],
                    res['kicker'],
                    res['total'], cur, delta]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                if i in (1, 2, 3, 9): style_data(c)
                elif i in (4, 7): style_int(c)
                elif i == 6: style_data(c); c.number_format = '0.0%'
                else: style_dollar(c)
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=9).fill = GATE_PASS_FILL if res['gate_met'] else GATE_MISS_FILL
            ws.cell(row=r, column=9).font = Font(bold=True)
            ws.cell(row=r, column=9).alignment = Alignment(horizontal='center')
            ws.cell(row=r, column=13).fill = CURRENT_FILL
            ws.cell(row=r, column=12).fill = PROP_A_FILL
            delta_cell = ws.cell(row=r, column=14)
            delta_cell.fill = GREEN_FILL if delta > 0 else (RED_FILL if delta < 0 else GREY_FILL)
            delta_cell.font = Font(bold=True, color='FFFFFF')
            delta_cell.alignment = Alignment(horizontal='center')
            r += 1

    # Footer total
    ws.cell(row=r, column=1, value=f'4-MONTH TOTAL - {label}').font = Font(bold=True)
    for col in range(1, 15): ws.cell(row=r, column=col).fill = SUB_FILL
    c = ws.cell(row=r, column=12, value=tot_prop); style_dollar(c); c.font = Font(bold=True)
    c = ws.cell(row=r, column=13, value=tot_cur);  style_dollar(c); c.font = Font(bold=True)
    c = ws.cell(row=r, column=14, value=tot_prop - tot_cur); style_dollar(c)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = GREEN_FILL if tot_prop - tot_cur > 0 else RED_FILL
    c.alignment = Alignment(horizontal='center')
    ws.row_dimensions[r].height = 26

    set_col_widths(ws, [22, 12, 9, 6, 10, 10, 6, 9, 9, 9, 8, 13, 10, 11])
    ws.freeze_panes = 'D5'


def build_rules(wb):
    ws = wb.create_sheet('Proposal Rules')
    ws['A1'] = "Boss Proposal Rules"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:F1')

    r = 3
    ws.cell(row=r, column=1, value='PRODUCER TIER LADDER').font = SECTION_FONT
    ws.cell(row=r, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    headers = ['Tier', 'Min NB count', 'Base salary', 'NB %', 'RN %', 'RWR %']
    for i, h in enumerate(headers, 1):
        style_header(ws.cell(row=r, column=i, value=h))
    r += 1
    for tier, t in PROPOSAL_TIERS.items():
        vals = [tier, t['min_nb_count'], t['base'], t['nb_pct'], t['rn_pct'], t['rwr_pct']]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            if i == 1: style_data(c); c.font = Font(bold=True)
            elif i in (2,): style_int(c)
            elif i == 3: style_dollar(c)
            else: style_data(c); c.number_format = '0.0%'
        r += 1
    r += 1

    ws.cell(row=r, column=1, value='SCENARIO DEFINITIONS').font = SECTION_FONT
    ws.cell(row=r, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    headers2 = ['Scenario', 'Description', '', '', '', '']
    for i, h in enumerate(headers2, 1):
        if h: style_header(ws.cell(row=r, column=i, value=h))
    r += 1
    for label, desc, _ in SCENARIOS:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        style_data(ws.cell(row=r, column=1))
        c = ws.cell(row=r, column=2, value=desc)
        style_data(c)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1
    r += 1

    ws.cell(row=r, column=1, value='OTHER RULES').font = SECTION_FONT
    ws.cell(row=r, column=1).fill = SECTION_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    rows = [
        ('$45K NB Premium Gate', 'Hard gate - variable + kicker paid ONLY if monthly NB written premium >= $45K.'),
        ('Carrier NB commission', '10% (used to compute NB agency commission)'),
        ('Carrier RN commission', '8% (used to compute RN agency commission)'),
        ('Carrier RWR commission', '8% (used to compute RWR agency commission)'),
        ('Collected kicker (NB only)', '<15% = $0 | 15-19% = +$1 | 20-24% = +$2 | 25-99% = +$5 | 100% PIF = +$8 per NB policy.'),
        ('Tier assignment', '4-month avg NB count drives tier per scenario (no sustained-month logic).'),
        ('NOT modeled', 'PIF / E-Pay / Ancillary / NSD bonuses (no per-agent monthly data).'),
    ]
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        style_data(ws.cell(row=r, column=1))
        c = ws.cell(row=r, column=2, value=v)
        style_data(c)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1

    set_col_widths(ws, [30, 18, 18, 14, 14, 14])


def main():
    scenario_results = [compute_scenario(sc[2]) for sc in SCENARIOS]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_scenarios_summary(wb, scenario_results)
    build_per_agent_grid(wb, scenario_results)
    build_tier_assignment_grid(wb, scenario_results)
    # One detail sheet per scenario
    sheet_names = [
        'Detail - REAL',
        'Detail - SWAP 50%',
        'Detail - FLIP 100%',
        'Detail - +25% NB + FLIP',
        'Detail - +50% NB + FLIP',
    ]
    for i, name in enumerate(sheet_names):
        build_month_detail(wb, scenario_results, i, name)
    build_rules(wb)

    out = '/home/user/fiesta-bonus/output/Boss_Proposal_Backtest.xlsx'
    wb.save(out)
    print(f"Saved: {out}")

    # Also print the headline summary to stdout
    print("\nHEADLINE - boss proposal under each scenario:")
    print(f"{'Scenario':<32} {'Current':>10} {'Proposal':>12} {'Delta':>11} {'Gate pass':>11}")
    for (label, _, _), result in zip(SCENARIOS, scenario_results):
        _, tot_cur, tot_prop, gate_pass = result
        delta = tot_prop - tot_cur
        print(f"{label:<32} ${tot_cur:>9,.0f} ${tot_prop:>11,.0f} ${delta:>+10,.0f} {gate_pass:>3}/24")


if __name__ == '__main__':
    main()
